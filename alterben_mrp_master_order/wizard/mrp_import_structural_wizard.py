# -*- coding: utf-8 -*-
from datetime import datetime, timedelta
import base64
import io
import re
import unicodedata

from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    import xlrd  # For .xls files
except Exception:  # pragma: no cover
    xlrd = None
try:
    import openpyxl  # For .xlsx files
except Exception:  # pragma: no cover
    openpyxl = None


class MrpImportStructuralWizard(models.TransientModel):
    _name = 'mrp.import.structural.wizard'
    _description = 'Importador de Estructural'

    file = fields.Binary(string="Archivo Excel", required=True)
    filename = fields.Char(string="Nombre de archivo")
    create_quotation = fields.Boolean(
        string="Subir cotizacion",
        default=False,
        help="Si esta activo, usa la hoja COTIZACION para crear una cotizacion en Odoo antes de crear las MOs.",
    )
    existing_sale_order_id = fields.Many2one(
        'sale.order',
        string="Cotizacion / SO existente",
        help="Obligatorio cuando no se crea la cotizacion desde el Excel.",
    )
    existing_sale_partner_id = fields.Many2one(
        'res.partner',
        string="Cliente",
        related='existing_sale_order_id.partner_id',
        readonly=True,
    )
    existing_sale_date_order = fields.Datetime(
        string="Fecha SO",
        related='existing_sale_order_id.date_order',
        readonly=True,
    )

    QUOTATION_HEADER_ALIASES = {
        'ruc': ['ruc'],
        'cliente': ['cliente'],
        'fecha': ['fecha'],
        'producto': ['producto', 'referencia'],
        'posicion': ['posicion', 'posición'],
        'piezas': ['piezas', 'pieza'],
        'largo': ['largo'],
        'alto': ['alto', 'ancho', 'anchura'],
        'cantidad': ['cantidad'],
        'precio': ['precio', 'pvp'],
    }
    FABRICATION_HEADER_ALIASES = {
        'componentes': ['componentes', 'componentes_producto'],
        'posicion': ['posicion', 'posición'],
        'por_consumir': ['por consumir', 'por_consumir'],
        'operacion': ['operacion', 'operaciones'],
        'duracion_esperada': ['duracion esperada', 'duracion_esperada'],
        'cantidad_empleados': ['cantidad empleados', 'cantidad_empleados'],
        'producto': ['producto', 'referencia'],
        'cantidad': ['cantidad', 'cantidad mo', 'can'],
        'origen': ['origen'],
        'pedido_original': ['pedido original', 'pedido_original'],
        'desechos': ['desechos', 'desecho', 'scrap'],
        'centro_trabajo': ['centro de trabajo', 'centro_trabajo'],
    }
    QUOTATION_REQUIRED_HEADERS = ['ruc', 'fecha', 'producto', 'piezas', 'largo', 'alto', 'precio']
    FABRICATION_REQUIRED_HEADERS = [
        'componentes', 'por_consumir', 'operacion', 'duracion_esperada', 'cantidad_empleados',
        'producto', 'posicion', 'cantidad', 'origen', 'pedido_original', 'desechos', 'centro_trabajo',
    ]

    @api.onchange('create_quotation')
    def _onchange_create_quotation(self):
        if self.create_quotation:
            self.existing_sale_order_id = False

    def _normalize_text(self, value):
        text = str(value or '').strip().lower()
        text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode('ascii')
        return re.sub(r'[^a-z0-9]', '', text)

    def _read_workbook(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Debes adjuntar un archivo para importar."))
        if not self.filename:
            raise UserError(_("Debes indicar el nombre del archivo."))

        content = base64.b64decode(self.file)
        fname = (self.filename or '').lower()
        if fname.endswith('.csv'):
            raise UserError(_("La importacion estructural ahora requiere un archivo Excel con hojas COTIZACION y FABRICACION."))

        sheets = {}
        if fname.endswith('.xlsx'):
            if openpyxl is None:
                raise UserError(_("No se puede leer Excel .xlsx porque falta la libreria Python 'openpyxl' en el servidor."))
            try:
                workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True, read_only=True)
            except Exception:
                raise UserError(_("No se pudo leer el archivo .xlsx. Verifica que este bien formado."))
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows_iter = sheet.iter_rows(values_only=True)
                try:
                    headers_row = next(rows_iter)
                except StopIteration:
                    sheets[sheet_name] = []
                    continue
                headers = [str(cell).strip() if cell is not None else '' for cell in headers_row]
                rows = []
                for row in rows_iter:
                    row_map = {}
                    for colx in range(len(headers)):
                        value = row[colx] if colx < len(row) else None
                        row_map[headers[colx]] = value
                    rows.append(row_map)
                sheets[sheet_name] = rows
            return sheets

        if xlrd is None:
            raise UserError(_("No se puede leer Excel .xls porque falta la libreria Python 'xlrd' en el servidor."))
        try:
            workbook = xlrd.open_workbook(file_contents=content)
        except Exception:
            raise UserError(_("No se pudo leer el archivo. Verifica que sea .xls o .xlsx."))
        for sheet in workbook.sheets():
            headers = [str(cell.value).strip() for cell in sheet.row(0)] if sheet.nrows else []
            rows = []
            for rowx in range(1, sheet.nrows):
                row_map = {}
                for colx in range(len(headers)):
                    row_map[headers[colx]] = sheet.cell(rowx, colx).value
                rows.append(row_map)
            sheets[sheet.name] = rows
        return sheets

    def _find_sheet(self, workbook_sheets, expected_name):
        expected = self._normalize_text(expected_name)
        for sheet_name, rows in workbook_sheets.items():
            normalized_sheet = self._normalize_text(sheet_name)
            if normalized_sheet == expected or normalized_sheet.startswith(expected):
                return sheet_name, rows
        return False, []

    def _resolve_headers(self, lines, aliases_map):
        if not lines:
            return {}
        raw_headers = list(lines[0].keys())
        header_map = {self._normalize_text(key): key for key in raw_headers}
        resolved = {}
        for canonical, aliases in aliases_map.items():
            for alias in aliases:
                normalized = self._normalize_text(alias)
                if normalized in header_map:
                    resolved[canonical] = header_map[normalized]
                    break
        return resolved

    def _get_value(self, row, header_map, canonical):
        key = header_map.get(canonical)
        return row.get(key) if key else None

    def _is_empty_row(self, row):
        return not any(str(value or '').strip() for value in row.values())

    def _to_float(self, value, row_num, column_name, errors):
        if value is None or value == '':
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(' ', '')
        if not text:
            return None
        if ',' in text and '.' in text:
            if text.rfind(',') > text.rfind('.'):
                text = text.replace('.', '').replace(',', '.')
            else:
                text = text.replace(',', '')
        else:
            text = text.replace(',', '.')
        try:
            return float(text)
        except Exception:
            errors.append(_("Fila %s, columna %s: valor invalido '%s'.") % (row_num, column_name, value))
            return None

    def _to_int(self, value, row_num, column_name, errors):
        number = self._to_float(value, row_num, column_name, errors)
        if number is None:
            return None
        if abs(number - round(number)) > 1e-6:
            errors.append(_("Fila %s, columna %s: debe ser entero, valor '%s'.") % (row_num, column_name, value))
            return None
        return int(round(number))

    def _to_datetime(self, value, row_num, column_name, errors):
        if value is None or value == '':
            return None
        if isinstance(value, datetime):
            return value
        text = str(value).strip()
        formats = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y")
        for fmt in formats:
            try:
                return datetime.strptime(text, fmt)
            except Exception:
                continue
        try:
            serial = float(text)
            if serial > 20000:
                return datetime(1899, 12, 30) + timedelta(days=serial)
        except Exception:
            pass
        errors.append(_("Fila %s, columna %s: fecha invalida '%s'.") % (row_num, column_name, value))
        return None

    def _extract_code(self, value):
        text = (value or '').strip()
        match = re.search(r'\[([^\]]+)\]', text)
        return match.group(1).strip() if match else text

    def _find_product(self, raw_value, row_num, column_name, errors):
        text = (raw_value or '').strip()
        if not text:
            errors.append(_("Fila %s, columna %s: valor vacio.") % (row_num, column_name))
            return False
        code = self._extract_code(text)
        Product = self.env['product.product']
        product = Product.search([('default_code', '=', code)], limit=2)
        if len(product) == 1:
            return product
        if len(product) > 1:
            errors.append(_("Fila %s, columna %s: codigo '%s' ambiguo.") % (row_num, column_name, code))
            return False
        product = Product.search([('name', 'ilike', text)], limit=2)
        if len(product) == 1:
            return product
        if len(product) > 1:
            errors.append(_("Fila %s, columna %s: nombre '%s' ambiguo.") % (row_num, column_name, text))
            return False
        errors.append(_("Fila %s, columna %s: producto no encontrado '%s'.") % (row_num, column_name, text))
        return False

    def _find_workcenter(self, name, row_num, errors):
        text = (name or '').strip()
        if not text:
            errors.append(_("Fila %s: Centro de trabajo vacio.") % row_num)
            return False
        alias_map = {
            'LAMINADO': 'AUTOCLAVE',
            'PREVACADO ESTRUCTURAL': 'PREVACIADO ESTRUCTURAL',
        }
        normalized = alias_map.get(text.upper(), text)
        Workcenter = self.env['mrp.workcenter']
        workcenter = Workcenter.search([('name', '=', normalized)], limit=2)
        if len(workcenter) == 1:
            return workcenter
        workcenter = Workcenter.search([('name', 'ilike', normalized)], limit=2)
        if len(workcenter) == 1:
            return workcenter
        errors.append(_("Fila %s: centro de trabajo no encontrado '%s'.") % (row_num, text))
        return False

    def _find_partner_by_ruc(self, ruc, row_num, errors):
        vat = str(ruc or "").strip()
        if not vat:
            errors.append(_("Fila %s: RUC vacio.") % row_num)
            return False
        Partner = self.env['res.partner'].with_context(active_test=False)
        partner = Partner.search([('vat', '=', vat)], limit=2)
        if len(partner) == 1:
            return partner
        errors.append(_("Fila %s: cliente con RUC '%s' no encontrado o ambiguo.") % (row_num, vat))
        return False

    def _get_structural_parameters(self):
        return self.env['mrp.master.type'].search([('active', '=', True)], order='id', limit=1)

    def _get_scrap_location(self):
        params = self._get_structural_parameters()
        if params and params.import_scrap_location_id:
            return params.import_scrap_location_id
        return self.env['stock.location'].search([('scrap_location', '=', True)], limit=1)

    def _parse_quotation_lines(self, lines):
        header_map = self._resolve_headers(lines, self.QUOTATION_HEADER_ALIASES)
        missing = [header for header in self.QUOTATION_REQUIRED_HEADERS if header not in header_map]
        if missing:
            raise UserError(_("Faltan columnas requeridas en la hoja COTIZACION: %s") % ", ".join(missing))

        errors = []
        parsed_lines = []
        partner = False
        date_order = False
        row_offset = 2
        for idx, row in enumerate(lines):
            row_num = idx + row_offset
            if self._is_empty_row(row):
                continue
            product_raw = str(self._get_value(row, header_map, 'producto') or '').strip()
            if not product_raw:
                continue
            product = self._find_product(product_raw, row_num, 'PRODUCTO', errors)
            row_partner = self._find_partner_by_ruc(self._get_value(row, header_map, 'ruc'), row_num, errors)
            row_date = self._to_datetime(self._get_value(row, header_map, 'fecha'), row_num, 'FECHA', errors)
            piezas = self._to_int(self._get_value(row, header_map, 'piezas'), row_num, 'PIEZAS', errors)
            largo = self._to_int(self._get_value(row, header_map, 'largo'), row_num, 'LARGO', errors)
            alto = self._to_int(self._get_value(row, header_map, 'alto'), row_num, 'ALTO', errors)
            qty_excel = self._to_float(self._get_value(row, header_map, 'cantidad'), row_num, 'CANTIDAD', errors)
            price_unit = self._to_float(self._get_value(row, header_map, 'precio'), row_num, 'PRECIO', errors)
            if not product or not row_partner or row_date is None or piezas is None or largo is None or alto is None or price_unit is None:
                continue
            qty_m2 = ((largo * alto) / 1000000.0) * piezas
            if qty_m2 <= 0:
                errors.append(_("Fila %s: la cantidad calculada en m2 debe ser mayor que cero.") % row_num)
                continue
            if partner and partner.id != row_partner.id:
                errors.append(_("Fila %s: la hoja COTIZACION contiene mas de un cliente.") % row_num)
                continue
            if date_order and date_order != row_date:
                errors.append(_("Fila %s: la hoja COTIZACION contiene mas de una fecha.") % row_num)
                continue
            partner = row_partner
            date_order = row_date
            parsed_lines.append({
                'row_num': row_num,
                'product': product,
                'qty_m2': qty_m2,
                'price_unit': price_unit,
                'piezas': piezas,
                'largo': largo,
                'alto': alto,
            })
        if errors:
            raise UserError("\n".join(errors))
        if not parsed_lines:
            raise UserError(_("La hoja COTIZACION no contiene lineas validas para importar."))
        return {
            'partner': partner,
            'date_order': date_order,
            'lines': parsed_lines,
        }

    def _parse_fabrication_groups(self, lines):
        header_map = self._resolve_headers(lines, self.FABRICATION_HEADER_ALIASES)
        missing = [header for header in self.FABRICATION_REQUIRED_HEADERS if header not in header_map]
        if missing:
            raise UserError(_("Faltan columnas requeridas en la hoja FABRICACION: %s") % ", ".join(missing))

        errors = []
        product_groups = {}
        row_offset = 2
        for idx, row in enumerate(lines):
            row_num = idx + row_offset
            if self._is_empty_row(row):
                continue
            product_raw = (self._get_value(row, header_map, 'producto') or '').strip()
            if not product_raw:
                errors.append(_("Fila %s: no hay PRODUCTO para asociar la fila.") % row_num)
                continue
            product_key = self._normalize_text(product_raw)
            current = product_groups.setdefault(product_key, {
                'row_start': row_num,
                'product_raw': product_raw,
                'qty_mo': 0.0,
                'qty_by_position': {},
                'pedido_original_values': [],
                'origins': [],
                'components': [],
                'operations': [],
                'operation_keys': set(),
            })
            current['row_start'] = min(current['row_start'], row_num)

            position_value = self._get_value(row, header_map, 'posicion')
            position_key = str(position_value or '').strip()
            qty_mo_raw = self._get_value(row, header_map, 'cantidad')
            qty_mo = self._to_float(qty_mo_raw, row_num, 'CANTIDAD', errors) if qty_mo_raw not in (None, '') else None
            if qty_mo is not None:
                if qty_mo <= 0:
                    errors.append(_("Fila %s: CANTIDAD debe ser mayor que cero.") % row_num)
                elif not position_key:
                    errors.append(_("Fila %s: POSICION vacia para la cantidad del producto.") % row_num)
                else:
                    previous_qty = current['qty_by_position'].get(position_key)
                    if previous_qty is None:
                        current['qty_by_position'][position_key] = qty_mo
                    elif abs(previous_qty - qty_mo) > 1e-6:
                        errors.append(_("Fila %s: CANTIDAD no coincide para la misma POSICION del producto.") % row_num)

            pedido_original = (self._get_value(row, header_map, 'pedido_original') or '').strip()
            if pedido_original:
                for pedido in [item.strip() for item in pedido_original.split('/') if item.strip()]:
                    if pedido not in current['pedido_original_values']:
                        current['pedido_original_values'].append(pedido)

            origin_value = (self._get_value(row, header_map, 'origen') or '').strip()
            if origin_value and origin_value not in current['origins']:
                current['origins'].append(origin_value)

            component_raw = (self._get_value(row, header_map, 'componentes') or '').strip()
            if component_raw:
                consume_qty = self._to_float(self._get_value(row, header_map, 'por_consumir'), row_num, 'POR CONSUMIR', errors)
                scrap_qty = self._to_float(self._get_value(row, header_map, 'desechos'), row_num, 'DESECHOS', errors)
                if consume_qty is None or consume_qty <= 0:
                    errors.append(_("Fila %s: POR CONSUMIR debe ser mayor que cero para el componente.") % row_num)
                else:
                    current['components'].append({
                        'raw': component_raw,
                        'qty': consume_qty,
                        'scrap_qty': scrap_qty or 0.0,
                        'row_num': row_num,
                    })

            operation_raw = (self._get_value(row, header_map, 'operacion') or '').strip()
            if operation_raw:
                workcenter_raw = (self._get_value(row, header_map, 'centro_trabajo') or '').strip()
                duration_expected = self._to_float(self._get_value(row, header_map, 'duracion_esperada'), row_num, 'DURACION ESPERADA', errors)
                employee_qty = self._to_int(self._get_value(row, header_map, 'cantidad_empleados'), row_num, 'CANTIDAD EMPLEADOS', errors)
                if not workcenter_raw:
                    errors.append(_("Fila %s: CENTRO DE TRABAJO vacio para la operacion.") % row_num)
                else:
                    op_key = (
                        self._normalize_text(operation_raw),
                        self._normalize_text(workcenter_raw),
                    )
                    if op_key not in current['operation_keys']:
                        current['operations'].append({
                            'name': operation_raw,
                            'workcenter': workcenter_raw,
                            'duration_expected': duration_expected,
                            'employee_qty': employee_qty,
                            'row_num': row_num,
                        })
                        current['operation_keys'].add(op_key)

        if errors:
            raise UserError("\n".join(errors))

        groups = []
        for group in product_groups.values():
            if not group['qty_by_position']:
                errors.append(_("Fila %s: no hay CANTIDAD valida por POSICION para '%s'.") % (group['row_start'], group['product_raw']))
                continue
            group['qty_mo'] = sum(group['qty_by_position'].values())
            group['pedido_original'] = "/".join(group['pedido_original_values']) if group['pedido_original_values'] else False
            group.pop('qty_by_position', None)
            group.pop('pedido_original_values', None)
            group.pop('operation_keys', None)
            groups.append(group)

        if errors:
            raise UserError("\n".join(errors))
        if not groups:
            raise UserError(_("La hoja FABRICACION no contiene bloques validos para importar."))
        return groups

    def _create_sale_order_from_quotation(self, quote_data, pedido_original_values):
        sale_order_model = self.env['sale.order']
        sale_line_model = self.env['sale.order.line']
        params = self._get_structural_parameters()

        so_vals = {
            'partner_id': quote_data['partner'].id,
            'date_order': fields.Datetime.to_string(quote_data['date_order']),
        }
        pedido_refs = [value for value in pedido_original_values if value]
        if pedido_refs and 'client_order_ref' in sale_order_model._fields:
            so_vals['client_order_ref'] = "/".join(pedido_refs)
        sale_order = sale_order_model.create(so_vals)

        for line in quote_data['lines']:
            line_vals = {
                'order_id': sale_order.id,
                'product_id': line['product'].id,
                'product_uom_qty': line['qty_m2'],
                'product_uom': line['product'].uom_id.id,
                'price_unit': line['price_unit'],
                'name': line['product'].display_name or line['product'].name,
            }
            if 'x_studio_largo' in sale_line_model._fields:
                line_vals['x_studio_largo'] = line['largo']
            if 'x_studio_ancho' in sale_line_model._fields:
                line_vals['x_studio_ancho'] = line['alto']
            if 'x_studio_piezas' in sale_line_model._fields:
                line_vals['x_studio_piezas'] = line['piezas']
            sale_line_model.create(line_vals)

        if params and params.structural_import_confirm_sale:
            sale_order.action_confirm()
        return sale_order

    def _find_reusable_bom(self, product):
        bom_model = self.env['mrp.bom']
        domain = ["|", ('product_id', '=', product.id), ('product_tmpl_id', '=', product.product_tmpl_id.id)]
        boms = bom_model.search(domain)
        if not boms:
            return bom_model.browse()

        def _is_macro_bom(bom):
            labels = [
                getattr(bom, 'code', False),
                getattr(bom, 'display_name', False),
                getattr(bom, 'product_tmpl_id', False) and bom.product_tmpl_id.display_name,
            ]
            haystack = " ".join([str(val or "") for val in labels]).upper()
            return "MACRO" in haystack

        macro_boms = boms.filtered(_is_macro_bom)
        if macro_boms:
            return macro_boms[:1]
        return boms[:1]

    def _create_bom_for_group(self, product, group, errors):
        bom_model = self.env['mrp.bom']
        bom_line_model = self.env['mrp.bom.line']
        routing_model = self.env['mrp.routing.workcenter']
        qty_mo = float(group.get('qty_mo') or 0.0)
        if qty_mo <= 0:
            errors.append(_("Fila %s: CANTIDAD de la MO debe ser mayor que cero para crear la BOM.") % (group['row_start']))
            return False, {}

        component_qty = {}
        component_details = {}
        for component in group['components']:
            comp_product = self._find_product(component['raw'], component['row_num'], 'COMPONENTES', errors)
            if not comp_product:
                continue
            component_qty[comp_product.id] = component_qty.get(comp_product.id, 0.0) + component['qty']
            detail = component_details.setdefault(comp_product.id, {'product': comp_product, 'scrap_qty': 0.0})
            detail['scrap_qty'] += component['scrap_qty']

        if not component_qty:
            errors.append(_("Fila %s: no hay componentes validos para '%s'.") % (group['row_start'], group['product_raw']))
            return False, {}

        bom_line_vals = []
        for prod_id, qty in component_qty.items():
            component_product = component_details[prod_id]['product']
            unit_qty = qty / qty_mo
            line_vals = {
                'product_id': prod_id,
                'product_qty': unit_qty,
            }
            if 'product_uom_id' in bom_line_model._fields:
                line_vals['product_uom_id'] = component_product.uom_id.id
            bom_line_vals.append((0, 0, line_vals))

        bom_vals = {
            'product_tmpl_id': product.product_tmpl_id.id,
            'product_id': product.id if 'product_id' in bom_model._fields else False,
            'product_uom_id': product.uom_id.id if 'product_uom_id' in bom_model._fields else False,
            'type': 'normal' if 'type' in bom_model._fields else False,
        }
        macro_label = "MACRO: %s" % (product.display_name or product.name or "")
        if 'code' in bom_model._fields:
            bom_vals['code'] = macro_label
        bom_vals = {key: value for key, value in bom_vals.items() if value is not False}
        existing_bom = self._find_reusable_bom(product)
        if existing_bom:
            bom = existing_bom
            if getattr(bom, 'bom_line_ids', False):
                bom.bom_line_ids.unlink()
            if 'bom_id' in routing_model._fields:
                routing_model.search([('bom_id', '=', bom.id)]).unlink()
            bom.write(dict(bom_vals, bom_line_ids=bom_line_vals))
        else:
            bom = bom_model.create(dict(bom_vals, bom_line_ids=bom_line_vals))
            if 'code' not in bom_model._fields and hasattr(bom, 'display_name'):
                try:
                    bom.name = macro_label
                except Exception:
                    pass

        seen_ops = set()
        for seq, operation in enumerate(group['operations'], start=1):
            workcenter = self._find_workcenter(operation['workcenter'], operation['row_num'], errors)
            if not workcenter:
                continue
            op_key = (operation['name'].strip().lower(), workcenter.id)
            if op_key in seen_ops:
                continue
            seen_ops.add(op_key)
            op_vals = {
                'name': operation['name'],
                'workcenter_id': workcenter.id,
            }
            if 'bom_id' in routing_model._fields:
                op_vals['bom_id'] = bom.id
            if 'sequence' in routing_model._fields:
                op_vals['sequence'] = seq
            if operation['duration_expected'] is not None:
                if 'time_cycle_manual' in routing_model._fields:
                    op_vals['time_cycle_manual'] = operation['duration_expected']
                elif 'duration_expected' in routing_model._fields:
                    op_vals['duration_expected'] = operation['duration_expected']
            routing_model.create(op_vals)

        return bom, component_details

    def _create_scraps_for_group(self, mo, group, component_details):
        params = self._get_structural_parameters()
        scrap_model = self.env['stock.scrap']
        scrap_location = self._get_scrap_location()
        created_scraps = self.env['stock.scrap']
        origin_ref = "/".join(group['origins']) if group['origins'] else (group['pedido_original'] or mo.origin or '')

        for detail in component_details.values():
            scrap_qty = detail['scrap_qty']
            if scrap_qty <= 0:
                continue
            product = detail['product']
            scrap_vals = {
                'product_id': product.id,
                'scrap_qty': scrap_qty,
                'product_uom_id': product.uom_id.id,
                'production_id': mo.id,
                'origin': origin_ref,
                'location_id': mo.location_src_id.id if getattr(mo, 'location_src_id', False) else False,
                'scrap_location_id': scrap_location.id if scrap_location else False,
            }
            if 'x_studio_pedido_original' in scrap_model._fields and group['pedido_original']:
                scrap_vals['x_studio_pedido_original'] = group['pedido_original']
            scrap_vals = {key: value for key, value in scrap_vals.items() if value is not False}
            scrap = scrap_model.create(scrap_vals)
            if params and params.structural_import_validate_scrap and getattr(scrap, 'state', False) == 'draft':
                scrap.action_validate()
            created_scraps |= scrap
        return created_scraps

    def action_import_structural(self):
        self.ensure_one()
        if not self.create_quotation and not self.existing_sale_order_id:
            raise UserError(_("Debe seleccionar una cotizacion/SO existente cuando no se active la opcion 'Subir cotizacion'."))

        workbook_sheets = self._read_workbook()
        quotation_sheet_name, quotation_rows = self._find_sheet(workbook_sheets, 'COTIZACION')
        fabrication_sheet_name, fabrication_rows = self._find_sheet(workbook_sheets, 'FABRICACION')
        if not fabrication_rows:
            raise UserError(_("No se encontro la hoja FABRICACION en el archivo."))

        fabrication_groups = self._parse_fabrication_groups(fabrication_rows)
        pedido_original_values = []
        for group in fabrication_groups:
            pedido = group.get('pedido_original')
            if pedido and pedido not in pedido_original_values:
                pedido_original_values.append(pedido)

        sale_order = self.existing_sale_order_id
        created_sale_ids = []
        if self.create_quotation:
            if not quotation_rows:
                raise UserError(_("No se encontro la hoja COTIZACION en el archivo."))
            quote_data = self._parse_quotation_lines(quotation_rows)
            sale_order = self._create_sale_order_from_quotation(quote_data, pedido_original_values)
            created_sale_ids.append(sale_order.id)

        params = self._get_structural_parameters()
        mrp_production = self.env['mrp.production']
        created_mo_ids = []
        created_bom_ids = []
        created_scrap_ids = []
        errors = []

        for group in fabrication_groups:
            if not group['product_raw']:
                continue
            if group['qty_mo'] is None:
                errors.append(_("Fila %s: CANTIDAD vacia para producto '%s'.") % (group['row_start'], group['product_raw']))
                continue
            product = self._find_product(group['product_raw'], group['row_start'], 'PRODUCTO', errors)
            if not product:
                continue

            bom, component_details = self._create_bom_for_group(product, group, errors)
            if not bom:
                continue
            created_bom_ids.append(bom.id)

            mo_vals = {
                'product_id': product.id,
                'product_qty': group['qty_mo'],
                'product_uom_id': product.uom_id.id if 'product_uom_id' in mrp_production._fields else False,
                'bom_id': bom.id,
                'origin': sale_order.name,
            }
            if 'x_studio_pedido_original' in mrp_production._fields and group['pedido_original']:
                mo_vals['x_studio_pedido_original'] = group['pedido_original']
            mo = mrp_production.create(mo_vals)
            created_mo_ids.append(mo.id)

            if params and params.structural_import_confirm_mo:
                try:
                    mo.action_confirm()
                except Exception as exc:
                    errors.append(_("Fila %s: error al confirmar MO %s: %s") % (group['row_start'], mo.display_name, str(exc)))

            created_scraps = self._create_scraps_for_group(mo, group, component_details)
            created_scrap_ids.extend(created_scraps.ids)

        summary = _(
            "Importacion completada. Cotizaciones: %s. MOs creadas: %s. BOMs creadas: %s. Desechos creados: %s."
        ) % (
            len(created_sale_ids),
            len(created_mo_ids),
            len(created_bom_ids),
            len(created_scrap_ids),
        )
        result = self.env['mrp.import.result.wizard'].create({
            'summary': summary,
            'error_message': "\n".join(errors) if errors else _("Sin errores."),
            'production_ids': [(6, 0, created_mo_ids)],
            'sale_order_ids': [(6, 0, created_sale_ids)],
            'scrap_ids': [(6, 0, created_scrap_ids)],
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'mrp.import.result.wizard',
            'view_mode': 'form',
            'res_id': result.id,
            'target': 'new',
        }
